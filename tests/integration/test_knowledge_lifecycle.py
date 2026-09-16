"""Reviewed local documents become bounded, scoped evidence in native ADK runs."""

import asyncio
import io
import json
import shutil
import sqlite3
from dataclasses import replace

from fastapi.testclient import TestClient
import pytest

from app.api.application import create_app
from app.configuration.knowledge import KnowledgeInput
from tests.support import FixtureModel, connectors, model_factory, settings_for


def draft(client, token, **changes):
    response = client.post("/api/v1/knowledge", headers=token("owner"), json={
        "title": "Checkout timeout runbook", "category": "Runbooks", "tags": ["checkout"],
        "content": "For checkout timeouts, inspect pool saturation and compare the affected customer journeys. Cite current incident observations before assigning a cause.",
    } | changes)
    assert response.status_code == 201, response.text
    return response.json()


def review(client, token, document, action, actor="owner", expected=None):
    return client.post(f"/api/v1/knowledge/{document['id']}/{action}", headers=token(actor), json={
        "expected_hash": expected or document["content_hash"], "reason": "Checked the exact document revision.",
    })


def approve(client, token, document):
    assert review(client, token, document, "submit").status_code == 200
    result = review(client, token, document, "approve", "admin")
    assert result.status_code == 200, result.text
    return result.json()


def test_knowledge_review_permissions_revisions_and_immutable_history(tmp_path):
    settings, token = settings_for(tmp_path)
    with TestClient(create_app(settings, connectors=connectors())) as client:
        item = draft(client, token)
        assert item["status"] == "draft" and not item["can_review"]
        assert client.get("/api/v1/knowledge", headers=token("viewer")).json() == []
        assert review(client, token, item, "submit").status_code == 200
        assert review(client, token, item, "approve").status_code == 403
        assert review(client, token, item, "approve", "viewer").status_code == 403
        assert review(client, token, item, "approve", "admin", "sha256:" + "0" * 64).status_code == 409
        assert review(client, token, item, "approve", "admin").status_code == 200
        assert client.get("/api/v1/knowledge", headers=token("viewer")).json()[0]["status"] == "approved"
        changed = client.put(f"/api/v1/knowledge/{item['id']}", headers=token("owner"), json={
            "title": item["title"], "content": "Updated checkout timeout guidance.", "expected_hash": item["content_hash"],
        })
        assert changed.status_code == 200, changed.text
        assert changed.json()["status"] == "draft" and changed.json()["revision"] == 2
        assert changed.json()["content_hash"] != item["content_hash"]
        assert review(client, token, item, "approve", "admin").status_code == 409
        history = client.get(f"/api/v1/knowledge/{item['id']}/history", headers=token("owner")).json()
        assert [event["action"] for event in history] == ["draft", "approve", "submit", "draft"]
        assert history[-1]["details"]["content_hash"] == item["content_hash"]
        saved = asyncio.run(client.app.state.knowledge.blobs.get(item["content_hash"]))
        assert json.loads(saved)["content"] == item["content"]
        for payload in ({"title": "Invalid", "content": "text", "status": "approved"},
                        {"title": "Invalid", "content": "text", "tenant_id": "outside"}):
            assert client.post("/api/v1/knowledge", headers=token("owner"), json=payload).status_code == 422


def test_approved_knowledge_is_cited_by_native_run_and_revocation_removes_future_use(tmp_path, monkeypatch):
    settings, token = settings_for(tmp_path, mode="live")
    seen = []
    original = FixtureModel.generate_content_async

    async def capture(self, request, stream=False):
        seen.append(str(request.config.system_instruction))
        async for response in original(self, request, stream):
            yield response

    monkeypatch.setattr(FixtureModel, "generate_content_async", capture)
    with TestClient(create_app(settings, connectors=connectors(), model_factory=model_factory)) as client:
        item = approve(client, token, draft(client, token))
        run = client.post("/api/v1/runs", headers=token("viewer"), json={
            "capability": "attachment_review", "prompt": "What should I check for checkout timeouts? Keep the answer short.",
        })
        assert run.status_code == 200, run.text
        assert run.json()["status"] == "SUCCEEDED", run.text
        run_id = run.json()["run_id"]
        evidence = client.get(f"/api/v1/runs/{run_id}/evidence", headers=token("viewer")).json()
        assert len(evidence) == 1 and evidence[0]["source"]["connector"] == "knowledge"
        assert json.loads(evidence[0]["query_json"])["doc_id"] == item["id"]
        assert evidence[0]["evidence_id"] in run.json()["result"]["findings"][0]["evidence_ids"]
        assert any("pool saturation" in instruction for instruction in seen)
        assert any("reference guidance, not a current observation" in instruction for instruction in seen)
        assert review(client, token, item, "revoke", "admin").status_code == 200
        next_run = client.post("/api/v1/runs", headers=token("viewer"), json={
            "capability": "attachment_review", "prompt": "What should I check for checkout timeouts?",
        })
        assert next_run.status_code == 200, next_run.text
        assert next_run.json()["evidence_count"] == 0
        assert client.get(f"/api/v1/runs/{run_id}/evidence", headers=token("viewer")).json() == evidence
    with sqlite3.connect(tmp_path / "runs.db") as database:
        contract = json.loads(database.execute("SELECT contract_json FROM runs WHERE run_id=?", (run_id,)).fetchone()[0])
        reference = json.loads(contract["model_config_json"])["knowledge_references"][0]
        assert reference["content_hash"] == item["content_hash"]
        assert reference["evidence_id"] == evidence[0]["evidence_id"]


def test_knowledge_retrieval_is_project_scoped_relevant_and_bounded(tmp_path):
    settings, token = settings_for(tmp_path)
    with TestClient(create_app(settings, connectors=connectors())) as client:
        item = approve(client, token, draft(client, token, content="checkout timeout " + "evidence " * 1000))
        service = client.app.state.knowledge
        principal = settings.principals["viewer"]
        selected = asyncio.run(service.relevant(principal, "checkout timeout", max_items=1, max_chars=200))
        assert len(selected) == 1 and selected[0]["doc_id"] == item["id"]
        assert len(selected[0]["excerpt"]) <= 200 and selected[0]["truncated"] is True
        assert asyncio.run(service.relevant(principal, "unrelated database", max_items=3, max_chars=1000)) == []
        other = principal.model_copy(update={"project_id": "other-project"})
        assert asyncio.run(service.relevant(other, "checkout timeout", max_items=3, max_chars=1000)) == []
        assert asyncio.run(service.list(other)) == []
        with pytest.raises(LookupError):
            asyncio.run(service.original(other, item["id"]))


def document_bytes(kind):
    output = io.BytesIO()
    if kind == "docx":
        from docx import Document
        doc = Document()
        doc.add_paragraph("Checkout timeout: inspect the connection pool.")
        doc.save(output)
    elif kind == "xlsx":
        from openpyxl import Workbook
        workbook = Workbook()
        workbook.active.append(["Service", "Investigation guidance"])
        workbook.active.append(["Checkout", "Inspect timeout evidence"])
        workbook.save(output)
    elif kind == "pdf":
        from pypdf import PdfWriter
        from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
        writer = PdfWriter()
        page = writer.add_blank_page(300, 200)
        font = DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"), NameObject("/BaseFont"): NameObject("/Helvetica")})
        page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})})
        stream = DecodedStreamObject()
        stream.set_data(b"BT /F1 14 Tf 20 100 Td (Checkout timeout guidance) Tj ET")
        page[NameObject("/Contents")] = writer._add_object(stream)
        writer.write(output)
    elif kind == "png":
        if not shutil.which("tesseract"):
            pytest.skip("Local Tesseract is unavailable")
        from PIL import Image, ImageDraw, ImageFont
        image = Image.new("RGB", (800, 160), "white")
        ImageDraw.Draw(image).text((30, 50), "Checkout timeout guidance", fill="black", font=ImageFont.load_default(size=38))
        image.save(output, format="PNG")
    else:
        output.write(b"service,guidance\ncheckout,inspect timeout evidence\n")
    return output.getvalue()


@pytest.mark.parametrize("kind", ["csv", "docx", "xlsx", "pdf", "png"])
def test_local_knowledge_file_formats_are_parsed_drafted_and_originals_retained(tmp_path, kind):
    settings, token = settings_for(tmp_path)
    data = document_bytes(kind)
    with TestClient(create_app(settings, connectors=connectors())) as client:
        uploaded = client.post("/api/v1/knowledge/upload", headers=token("owner"),
            data={"title": "Checkout guidance", "tags": '["checkout"]'},
            files={"file": (f"guidance.{kind}", data, "application/octet-stream")})
        assert uploaded.status_code == 201, uploaded.text
        item = uploaded.json()
        assert item["status"] == "draft" and "checkout" in item["content"].lower()
        assert item["tags"] == ["checkout"] and item["upload"]["original_retained"] is True
        original = client.get(f"/api/v1/knowledge/{item['id']}/download", headers=token("owner"))
        assert original.status_code == 200 and original.content == data
        assert client.get(f"/api/v1/knowledge/{item['id']}/download", headers=token("viewer")).status_code == 403
        replacement = client.post("/api/v1/knowledge/upload", headers=token("owner"),
            data={"title": "Checkout updated", "doc_id": item["id"], "expected_hash": item["content_hash"]},
            files={"file": ("replacement.txt", b"Updated checkout timeout guidance.", "text/plain")})
        assert replacement.status_code == 201, replacement.text
        assert replacement.json()["revision"] == 2


def test_knowledge_upload_uses_active_file_limits_and_rejects_untrusted_fields(tmp_path):
    settings, token = settings_for(tmp_path)
    with TestClient(create_app(settings, connectors=connectors())) as client:
        client.app.state.file_limits = replace(client.app.state.file_limits, max_file_bytes=8)
        response = client.post("/api/v1/knowledge/upload", headers=token("owner"),
            data={"title": "Too large"}, files={"file": ("large.txt", b"x" * 9, "text/plain")})
        assert response.status_code == 413
        assert client.post("/api/v1/knowledge/upload", data={"title": "Unauthenticated"},
            files={"file": ("x.txt", b"x", "text/plain")}).status_code == 401
        with pytest.raises(ValueError):
            KnowledgeInput(title="", content="text")
